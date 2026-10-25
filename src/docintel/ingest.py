"""Turn an uploaded file into pages: an optional embedded text layer plus a rendered image."""
import io
from pathlib import PurePath
from PIL import Image, ImageOps
from dataclasses import dataclass, field

SUPPORTED = {".pdf", ".docx", ".png", ".jpg", ".jpeg", ".tif", ".tiff"}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".tif", ".tiff"}
MAX_PAGES = 200
RENDER_SCALE = 2.0  # ~144 dpi: good OCR accuracy at reasonable speed
MIN_TEXT_LAYER_CHARS = 40


class UnsupportedDocument(ValueError):
    pass


@dataclass
class Page:
    page_no: int
    text_layer: str = ""
    image: Image.Image | None = field(default=None, repr=False)

    @property
    def has_text_layer(self) -> bool:
        return len(self.text_layer.strip()) >= MIN_TEXT_LAYER_CHARS


def load(filename: str, data: bytes, max_bytes: int = 50 * 1024 * 1024) -> list[Page]:
    ext = PurePath(filename).suffix.lower()
    if ext not in SUPPORTED:
        raise UnsupportedDocument(f"Unsupported file type '{ext}'. Supported: {', '.join(sorted(SUPPORTED))}")
    if len(data) > max_bytes:
        raise UnsupportedDocument(f"File is larger than {max_bytes // (1024 * 1024)} MB")
    if not data:
        raise UnsupportedDocument("File is empty")
    try:
        if ext == ".pdf":
            return _load_pdf(data)
        if ext == ".docx":
            return _load_docx(data)
        return _load_image(data)
    except UnsupportedDocument:
        raise
    except Exception as e:  # corrupt files raise many library-specific exception types
        raise UnsupportedDocument(f"Could not read {ext} file: {e}") from e


def _load_pdf(data: bytes) -> list[Page]:
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(data)
    if len(pdf) > MAX_PAGES:
        raise UnsupportedDocument(f"PDF has {len(pdf)} pages; the limit is {MAX_PAGES}")
    pages = []
    for i in range(len(pdf)):
        page = pdf[i]
        text = page.get_textpage().get_text_bounded() or ""
        img = page.render(scale=RENDER_SCALE).to_pil().convert("RGB")
        pages.append(Page(i + 1, text.replace("\r\n", "\n"), img))
    return pages


def _load_docx(data: bytes) -> list[Page]:
    import docx

    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    text = "\n".join(parts)
    if not text.strip():
        raise UnsupportedDocument("DOCX contains no text")
    return [Page(1, text, None)]


def _load_image(data: bytes) -> list[Page]:
    img = Image.open(io.BytesIO(data))
    pages = []
    for i in range(getattr(img, "n_frames", 1)):
        img.seek(i)
        frame = ImageOps.exif_transpose(img.copy()).convert("RGB")
        pages.append(Page(i + 1, "", frame))
        if len(pages) >= MAX_PAGES:
            break
    return pages
